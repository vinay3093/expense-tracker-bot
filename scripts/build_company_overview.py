#!/usr/bin/env python3
"""Build docs/company/Expense_Tracker_Platform_Overview.docx for company sharing."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "company" / "Expense_Tracker_Platform_Overview.docx"
DIAGRAM = ROOT / "docs" / "company" / "architecture-diagram.mmd"


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for p in t.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            t.rows[ri + 1].cells[ci].text = cell


def build() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Cover ─────────────────────────────────────────────────────────────
    title = doc.add_heading("Conversational Expense Tracker Platform", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Enterprise Architecture & Capability Overview\n")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)
    sub.add_run(f"Document date: {date.today():%B %d, %Y}\n")
    sub.add_run("Classification: Internal — Technical Overview\n")
    sub.add_run("Repository: expense-tracker-bot")

    doc.add_page_break()

    # ── 1 Executive Summary ───────────────────────────────────────────────
    _heading(doc, "1. Executive Summary")
    doc.add_paragraph(
        "The Conversational Expense Tracker Platform is a production-grade, "
        "chat-driven personal finance assistant. Authorized users interact in "
        "natural language via Telegram; the system interprets intent, extracts "
        "structured expense data using a large language model (LLM), persists "
        "records to a configurable ledger, and answers analytical questions "
        "such as category totals, date-range spend, and period-over-period "
        "comparisons."
    )
    doc.add_paragraph(
        "The solution is designed with enterprise software patterns—layered "
        "architecture, provider abstraction, typed configuration, comprehensive "
        "automated testing, and observable audit trails—while remaining "
        "lightweight to deploy on free-tier or self-hosted infrastructure."
    )

    _heading(doc, "1.1 Business Capabilities", level=2)
    _bullets(
        doc,
        [
            "Log expenses from conversational messages (e.g., “spent $42 on groceries today”).",
            "Query historical spend by category, date, or natural-language filters.",
            "Generate summaries (week, month, year) with optional comparisons to prior periods.",
            "Correct entries via undo and edit commands with ledger consistency.",
            "Multi-currency support with automatic foreign-exchange conversion where configured.",
        ],
    )

    _heading(doc, "1.2 Design Principles", level=2)
    _bullets(
        doc,
        [
            "Secure by default — explicit user allow-list; no anonymous public access.",
            "Auditable — every chat turn and LLM invocation is traceable in append-only logs.",
            "Storage-agnostic — one ledger interface; switch Google Sheets or Postgres with a single setting.",
            "Provider-agnostic LLM — Groq, OpenAI, Anthropic, or local Ollama without rewriting business logic.",
            "Fail gracefully — user-facing errors are plain language; stack traces never surface in chat.",
        ],
    )

    doc.add_page_break()

    # ── 2 Architecture ────────────────────────────────────────────────────
    _heading(doc, "2. System Architecture")
    doc.add_paragraph(
        "The platform follows a five-layer architecture. Each layer exposes "
        "narrow interfaces so that presentation, orchestration, intelligence, "
        "and persistence can evolve independently."
    )

    _table(
        doc,
        ["Layer", "Responsibility", "Key modules"],
        [
            [
                "Presentation",
                "Telegram gateway, command handlers, health checks, access control",
                "telegram_app/",
            ],
            [
                "Orchestration",
                "End-to-end chat pipeline: classify → act → reply; write/read/correct paths",
                "pipeline/",
            ],
            [
                "Intelligence",
                "Intent classification, structured JSON extraction, category vocabulary",
                "extractor/ · llm/",
            ],
            [
                "Persistence",
                "Ledger abstraction; Sheets or Postgres backends; FX and workbook builders",
                "ledger/",
            ],
            [
                "Observability",
                "Conversation history and LLM call traces for debugging and audit",
                "storage/",
            ],
        ],
    )

    _heading(doc, "2.1 Architecture Diagram", level=2)
    diagram_png = ROOT / "docs" / "company" / "architecture-diagram.png"
    if diagram_png.is_file():
        doc.add_picture(str(diagram_png), width=Inches(6.5))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run("Figure 1 — Platform architecture (high level)").italic = True
        doc.add_paragraph()
    doc.add_paragraph(
        "An editable Mermaid source is maintained at "
        "docs/company/architecture-diagram.mmd for Confluence, GitHub, or "
        "presentation exports (PNG/SVG via mermaid.live)."
    )

    _heading(doc, "2.2 Request Flow (Log Expense)", level=2)
    flow_steps = [
        "User sends a natural-language message on Telegram.",
        "Presentation layer validates the sender against the configured allow-list.",
        "Orchestration invokes the extractor to classify intent and parse structured fields.",
        "LLM abstraction calls the configured provider (JSON-mode, temperature-controlled).",
        "Write path maps extracted fields to ledger rows (category, amount, date, currency).",
        "Optional FX conversion normalizes amounts to a base currency.",
        "Ledger backend persists the transaction (Google Sheets tab or Postgres row).",
        "A conversational confirmation is returned; the full turn is logged for audit.",
    ]
    for i, step in enumerate(flow_steps, 1):
        doc.add_paragraph(f"{i}. {step}", style="List Number")

    _heading(doc, "2.3 Request Flow (Query / Summary)", level=2)
    doc.add_paragraph(
        "For query intents, the orchestration layer routes to the retrieval engine, "
        "which reads from the active ledger backend, applies filters derived from "
        "the user’s question, aggregates results, and formats a natural-language "
        "reply. No write operations occur on the query path."
    )

    doc.add_page_break()

    # ── 3 Components ──────────────────────────────────────────────────────
    _heading(doc, "3. Component Overview")

    _heading(doc, "3.1 Telegram Front-End", level=2)
    doc.add_paragraph(
        "Built on the python-telegram-bot SDK. Handles inbound messages, slash "
        "commands (/undo, /edit), and exposes a minimal HTTP /health endpoint "
        "for platform health checks (Render, systemd, Docker)."
    )

    _heading(doc, "3.2 Conversation Pipeline", level=2)
    doc.add_paragraph(
        "Central orchestrator composing extraction, logging, retrieval, summary, "
        "correction, and reply formatting. Ensures one user message produces one "
        "persisted conversation turn with complete action metadata."
    )

    _heading(doc, "3.3 LLM & Extraction", level=2)
    doc.add_paragraph(
        "Multi-provider LLM factory with retries, timeouts, JSON repair, and "
        "structured tracing. Default production configuration uses Groq Cloud "
        "(openai/gpt-oss-20b) on the free tier; providers are swappable via "
        "environment configuration without code changes."
    )

    _heading(doc, "3.4 Ledger Backends", level=2)
    _table(
        doc,
        ["Edition", "Use case", "Technology"],
        [
            [
                "Google Sheets (default)",
                "Zero infrastructure; mobile-friendly spreadsheet view",
                "Google Sheets API · service account",
            ],
            [
                "Postgres / NocoDB",
                "Typed schema, audit log, soft-delete, long-term scale",
                "PostgreSQL · SQLAlchemy · Alembic · NocoDB UI",
            ],
            [
                "Mirror (roadmap)",
                "Sheets for daily use + Postgres for durable backup/analytics",
                "Dual-write reconciliation",
            ],
        ],
    )

    doc.add_page_break()

    # ── 4 Security ────────────────────────────────────────────────────────
    _heading(doc, "4. Security & Governance")
    _table(
        doc,
        ["Control", "Implementation"],
        [
            ["Authentication", "Telegram user ID allow-list (deny-all when unset)"],
            ["Secrets", "Environment variables / .env — never committed to source control"],
            ["Least privilege", "Google service account scoped to one spreadsheet; DB role limited to app schema"],
            ["Data isolation", "Single-tenant personal ledger per deployment instance"],
            ["Observability", "JSONL logs for conversations and LLM calls (local or mounted volume)"],
            ["Failure handling", "Errors contained in pipeline; no sensitive data in user replies"],
        ],
    )

    # ── 5 Deployment ────────────────────────────────────────────────────────
    _heading(doc, "5. Deployment Options")
    _table(
        doc,
        ["Target", "Profile", "Notes"],
        [
            ["Render Free", "Managed Docker, 24/7 with keep-alive", "Recommended starter; Sheets edition"],
            ["Oracle Cloud Free VM", "systemd self-hosted", "Full control; Sheets or NocoDB stack"],
            ["Docker", "Portable container", "Multi-stage Dockerfile included"],
            ["Local CLI", "Development / offline", "Ollama provider; fake LLM for tests"],
        ],
    )
    doc.add_paragraph(
        "Continuous integration runs automated tests (Python 3.10–3.12), linting, "
        "and Dockerfile validation on every push via GitHub Actions."
    )

    # ── 6 Technology Stack ────────────────────────────────────────────────
    _heading(doc, "6. Technology Stack")
    _table(
        doc,
        ["Category", "Technologies"],
        [
            ["Language", "Python 3.10+"],
            ["Chat interface", "Telegram Bot API · python-telegram-bot"],
            ["LLM integration", "Groq SDK · OpenAI-compatible clients · LiteLLM-ready patterns"],
            ["Sheets ledger", "gspread · Google service account"],
            ["Postgres ledger", "SQLAlchemy 2 · Alembic · NocoDB"],
            ["Configuration", "pydantic-settings (typed, validated env)"],
            ["Testing", "pytest (~500+ hermetic unit tests)"],
            ["Packaging", "pyproject.toml · Docker multi-stage build"],
        ],
    )

    # ── 7 Quality ─────────────────────────────────────────────────────────
    _heading(doc, "7. Quality Assurance")
    _bullets(
        doc,
        [
            "Hermetic unit tests with fake LLM and in-memory ledger backends.",
            "Module-level test coverage across extractor, pipeline, ledger, and Telegram auth.",
            "Structured JSON schemas validated on every LLM extraction response.",
            "Self-healing sheet format rebuilders for monthly/year/YTD workbook consistency.",
        ],
    )

    # ── 8 Roadmap ───────────────────────────────────────────────────────────
    _heading(doc, "8. Evolution & Roadmap")
    _bullets(
        doc,
        [
            "Mirror storage edition — dual-write Sheets + Postgres with reconciliation.",
            "Additional LLM providers and cost-aware routing.",
            "Enhanced analytics exports from Postgres ledger.",
        ],
    )

    # ── Appendix ──────────────────────────────────────────────────────────
    doc.add_page_break()
    _heading(doc, "Appendix A — Diagram Source")
    doc.add_paragraph(
        f"Mermaid source file: {DIAGRAM.relative_to(ROOT).as_posix()}\n\n"
        "To export for presentations:\n"
        "1. Open https://mermaid.live\n"
        "2. Paste the contents of architecture-diagram.mmd\n"
        "3. Export as PNG or SVG\n"
        "4. Insert into PowerPoint or this document as needed."
    )

    _heading(doc, "Appendix B — Related Documentation")
    _bullets(
        doc,
        [
            "docs/HANDBOOK.md — Full operational setup and module reference",
            "docs/ARCHITECTURE.md — Developer module map",
            "docs/MASTER_GUIDE.md — Deployment history and runbooks",
            "README.md — Repository entry point",
        ],
    )

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run(
        "— End of document — Generated for internal technical communication. "
        "Operational details may change; refer to the repository for the latest configuration."
    ).italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
